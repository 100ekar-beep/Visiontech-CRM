import datetime
import hashlib
import html
import hmac
import io
import math
import re
from urllib.parse import urlparse

import pandas as pd
import streamlit as st
from supabase import Client, create_client


st.set_page_config(page_title="Banking", page_icon="🏦", layout="wide")

ALLOWED_WORKSPACE = "VISPL"
PAGE_SIZE = 20
PAGE_PASSWORD_HASH = "7559103d4a3e65bbdc120e9e0bf6ad3c542690213b4d74e29ddb4b51ea00a261"

ACCOUNTS = {
    "HDFC VISPL": {
        "key": "HDFC_VISPL",
        "pay_from": "VISPL HDFC",
    },
    "HDFC VIS": {
        "key": "HDFC_VIS",
        "pay_from": "Visiontech Partnership",
    },
    "Pramodkumar Jaju": {
        "key": "PRAMODKUMAR_JAJU",
        "pay_from": "Pramodkumar Jaju",
    },
    "Radhika Jaju": {
        "key": "RADHIKA_JAJU",
        "pay_from": "Radhika Jaju",
    },
}

MATERIAL_VENDORS = {
    "RUSHIKESH STEEL": {
        "key": "MATERIAL_RUSHIKESH_STEEL",
        "pay_from": "RUSHIKESH STEEL",
    },
    "RUSHIKESH STEEL CORPORATION": {
        "key": "MATERIAL_RUSHIKESH_STEEL_CORPORATION",
        "pay_from": "RUSHIKESH STEEL CORPORATION",
    },
    "Vighanharta Enterprises": {
        "key": "MATERIAL_VIGHANHARTA_ENTERPRISES",
        "pay_from": "Vighanharta Enterprises",
    },
    "S M ENTERPRISES": {
        "key": "MATERIAL_S_M_ENTERPRISES",
        "pay_from": "S M ENTERPRISES",
    },
    "S K AND SONS ENTERPRISES": {
        "key": "MATERIAL_S_K_AND_SONS_ENTERPRISES",
        "pay_from": "S K AND SONS ENTERPRISES",
    },
}

MAIN_ACCOUNT_TABS = list(ACCOUNTS.keys()) + ["Material"]


st.markdown(
    """
    <style>
    .stApp { background: linear-gradient(135deg, #f8fafc 0%, #eef2ff 100%); }

    /* Premium sidebar navigation — same as Site Data page */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #0f172a 0%, #1e1b4b 100%);
        border-right: 1px solid rgba(255, 255, 255, 0.05);
    }
    [data-testid="stSidebarNav"] a {
        padding: 0.85rem 1.2rem !important;
        margin: 0.5rem 1rem !important;
        border-radius: 12px !important;
        background: rgba(255, 255, 255, 0.03) !important;
        color: #cbd5e1 !important;
        font-weight: 600 !important;
        font-size: 1.05rem !important;
        transition: all 0.3s ease !important;
        border: 1px solid rgba(255, 255, 255, 0.05) !important;
        display: flex !important;
        align-items: center !important;
        gap: 12px !important;
    }
    [data-testid="stSidebarNav"] a:hover {
        background: rgba(255, 255, 255, 0.1) !important;
        transform: translateX(4px) !important;
        border-color: rgba(255, 255, 255, 0.2) !important;
        color: #ffffff !important;
    }
    [data-testid="stSidebarNav"] a[aria-current="page"] {
        background: linear-gradient(90deg, #3b82f6 0%, #8b5cf6 100%) !important;
        color: #ffffff !important;
        border-color: transparent !important;
        box-shadow: 0 4px 15px rgba(59, 130, 246, 0.4) !important;
    }
    [data-testid="stSidebarNav"] a span {
        color: inherit !important;
    }
    div.stButton > button {
        background: linear-gradient(90deg, #3b82f6 0%, #8b5cf6 100%);
        color: white !important; border: none; border-radius: 10px;
        font-weight: 800 !important; padding: 0.65rem 1rem;
        min-height: 48px; transition: all 0.25s ease;
        box-shadow: 0 5px 12px rgba(59, 130, 246, 0.25);
    }
    div.stButton > button:hover {
        transform: translateY(-2px);
        box-shadow: 0 10px 20px rgba(79, 70, 229, 0.32);
    }
    div.stButton > button:disabled {
        background: #cbd5e1 !important; color: #64748b !important;
        box-shadow: none !important; transform: none !important;
        cursor: not-allowed !important; opacity: 0.75 !important;
    }
    div.stButton > button:disabled p,
    div.stButton > button:disabled span,
    div.stButton > button:disabled div { color: #64748b !important; }
    div.stButton > button p,
    div.stButton > button span,
    div.stButton > button div { color: #ffffff !important; font-weight: 800 !important; }
    div.stDownloadButton > button {
        background: linear-gradient(90deg, #059669 0%, #0d9488 100%) !important;
        color: #ffffff !important; border: none !important; border-radius: 10px !important;
        font-weight: 800 !important; min-height: 48px !important;
        box-shadow: 0 5px 12px rgba(5,150,105,0.24) !important;
        transition: all 0.25s ease !important;
    }
    div.stDownloadButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 10px 20px rgba(5,150,105,0.30) !important;
    }
    div.stDownloadButton > button p,
    div.stDownloadButton > button span,
    div.stDownloadButton > button div { color: #ffffff !important; font-weight: 800 !important; }

    .st-key-banking_account_nav div[data-testid="stHorizontalBlock"],
    .st-key-banking_view_nav div[data-testid="stHorizontalBlock"],
    .st-key-banking_material_vendor_nav div[data-testid="stHorizontalBlock"] {
        gap: 14px !important; flex-wrap: wrap !important;
    }
    .st-key-banking_account_nav button,
    .st-key-banking_view_nav button,
    .st-key-banking_material_vendor_nav button {
        font-size: 1.05rem !important; font-weight: 800 !important;
        min-height: 58px !important; padding: 14px 12px !important;
        border-radius: 13px !important; white-space: nowrap !important;
        transition: all 0.25s ease !important;
    }
    .st-key-banking_account_nav button[kind="secondary"],
    .st-key-banking_view_nav button[kind="secondary"],
    .st-key-banking_material_vendor_nav button[kind="secondary"] {
        background: #ffffff !important; color: #475569 !important;
        border: 1.5px solid rgba(15,23,42,0.12) !important;
        box-shadow: 0 3px 8px rgba(15,23,42,0.08) !important;
    }
    .st-key-banking_account_nav button[kind="secondary"] p,
    .st-key-banking_account_nav button[kind="secondary"] span,
    .st-key-banking_account_nav button[kind="secondary"] div,
    .st-key-banking_view_nav button[kind="secondary"] p,
    .st-key-banking_view_nav button[kind="secondary"] span,
    .st-key-banking_view_nav button[kind="secondary"] div,
    .st-key-banking_material_vendor_nav button[kind="secondary"] p,
    .st-key-banking_material_vendor_nav button[kind="secondary"] span,
    .st-key-banking_material_vendor_nav button[kind="secondary"] div {
        color: #475569 !important; font-weight: 800 !important;
    }
    .st-key-banking_account_nav button[kind="secondary"]:hover,
    .st-key-banking_view_nav button[kind="secondary"]:hover,
    .st-key-banking_material_vendor_nav button[kind="secondary"]:hover {
        background: #f8fafc !important; transform: translateY(-2px) !important;
        border-color: rgba(79,70,229,0.35) !important;
    }
    .st-key-banking_account_nav button[kind="primary"],
    .st-key-banking_view_nav button[kind="primary"],
    .st-key-banking_material_vendor_nav button[kind="primary"] {
        background: linear-gradient(90deg, #3b82f6 0%, #8b5cf6 100%) !important;
        color: #ffffff !important; border: none !important;
        box-shadow: 0 7px 18px rgba(79,70,229,0.38) !important;
    }
    .bank-title {
        color: white; padding: 28px 24px; border-radius: 14px;
        background: linear-gradient(90deg, #3b82f6 0%, #8b5cf6 52%, #ec4899 100%);
        box-shadow: 0 9px 24px rgba(79,70,229,.30); margin: 14px 0 22px;
        text-align: center;
    }
    .bank-title h1 { color: white !important; margin: 0; font-size: 2.35rem; letter-spacing: 1px; }
    .bank-title p { color: #e0e7ff !important; margin: 5px 0 0; font-weight: 600; }
    .table-head {
        background: #312e81; color: white; border-radius: 9px;
        padding: 10px 8px; font-weight: 800; margin-top: 8px;
    }
    .txn-row {
        background: white; border: 1px solid #e2e8f0; border-radius: 10px;
        padding: 9px 8px; margin: 6px 0; box-shadow: 0 2px 7px rgba(15,23,42,.05);
    }
    .narration { white-space: normal; overflow-wrap: anywhere; line-height: 1.35; }
    .amount { color: #dc2626; font-weight: 900; font-size: 1.02rem; }
    .approved { color: #047857; font-weight: 800; }
    .preview-wrap {
        max-height: 480px; overflow: auto; background: white;
        border: 1px solid #dbe3ef; border-radius: 12px;
    }
    .preview-table { width: 100%; border-collapse: collapse; table-layout: fixed; }
    .preview-table th {
        position: sticky; top: 0; z-index: 1; background: #312e81; color: white;
        padding: 11px 9px; text-align: left; font-weight: 800;
    }
    .preview-table td {
        padding: 9px; border-bottom: 1px solid #e5e7eb; vertical-align: top;
        color: #1f2937;
    }
    .preview-table .p-date { width: 125px; white-space: nowrap; }
    .preview-table .p-narration { width: auto; white-space: normal; overflow-wrap: anywhere; }
    .preview-table .p-ref { width: 205px; overflow-wrap: anywhere; }
    .preview-table .p-amount { width: 125px; white-space: nowrap; text-align: right; font-weight: 800; }
    .bulk-panel-title {
        background: linear-gradient(90deg, #2563eb 0%, #7c3aed 55%, #db2777 100%);
        color: #ffffff; border-radius: 13px; padding: 16px 20px;
        font-size: 1.15rem; font-weight: 900; letter-spacing: 0.3px;
        box-shadow: 0 7px 18px rgba(79,70,229,0.28); margin-bottom: 14px;
    }
    .st-key-bulk_move_panel [data-testid="stVerticalBlockBorderWrapper"] {
        background: rgba(255,255,255,0.94) !important;
        border: 1px solid rgba(99,102,241,0.20) !important;
        border-radius: 16px !important; padding: 8px !important;
        box-shadow: 0 8px 24px rgba(15,23,42,0.09) !important;
    }
    .st-key-bulk_move_panel input {
        min-height: 52px !important; border-radius: 11px !important;
        font-size: 1rem !important; font-weight: 650 !important;
    }
    div[data-baseweb="select"] * { font-weight: 700 !important; }
    </style>
    """,
    unsafe_allow_html=True,
)


# Always keep a visible way back, including on password/error screens.
if st.button("← Back to Home", key="banking_back_home", use_container_width=False):
    st.switch_page("app.py")


active_workspace = st.session_state.get("active_workspace", "VISPL")
if active_workspace != ALLOWED_WORKSPACE:
    st.error("Access Restricted")
    st.warning("Banking module केवल VISPL login के लिए उपलब्ध है।")
    st.stop()


# Separate password protection for this financial page.
if not st.session_state.get("banking_page_unlocked", False):
    st.markdown(
        """
        <div class="bank-title">
            <h1>Banking Access</h1>
            <p>इस financial page को खोलने के लिए personal password डालें।</p>
        </div>
        """,
        unsafe_allow_html=True,
    )
    with st.form("banking_password_form", clear_on_submit=True):
        entered_password = st.text_input("Personal Password", type="password")
        unlock_clicked = st.form_submit_button(
            "Unlock Banking",
            type="primary",
            use_container_width=True,
        )

    if unlock_clicked:
        entered_hash = hashlib.sha256(entered_password.encode("utf-8")).hexdigest()
        if hmac.compare_digest(entered_hash, PAGE_PASSWORD_HASH):
            st.session_state["banking_page_unlocked"] = True
            st.rerun()
        else:
            st.error("Incorrect Banking password.")
    st.stop()

with st.sidebar:
    if st.button("Lock Banking", use_container_width=True):
        st.session_state["banking_page_unlocked"] = False
        st.rerun()


@st.cache_resource
def init_connection() -> Client:
    supabase_section = st.secrets.get("supabase", {})

    raw_url = (
        supabase_section.get("url")
        or supabase_section.get("SUPABASE_URL")
        or st.secrets.get("SUPABASE_URL")
        or st.secrets.get("supabase_url")
    )
    raw_key = (
        supabase_section.get("key")
        or supabase_section.get("SUPABASE_KEY")
        or st.secrets.get("SUPABASE_KEY")
        or st.secrets.get("supabase_key")
    )

    if not raw_url or not raw_key:
        raise ValueError("Supabase URL/Key is missing in Streamlit Secrets")

    url = str(raw_url).strip().strip('"').strip("'").strip()
    key = str(raw_key).strip().strip('"').strip("'").strip()

    # create_client needs the project base URL, not the REST endpoint URL.
    if "/rest/v1" in url:
        url = url.split("/rest/v1", 1)[0]
    url = url.rstrip("/")
    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    parsed_url = urlparse(url)
    if not parsed_url.scheme or not parsed_url.netloc:
        raise ValueError("Invalid Supabase project URL in Streamlit Secrets")

    return create_client(url, key)


try:
    supabase: Client = init_connection()
except Exception as exc:
    st.error(f"Supabase connection error: {exc}")
    st.stop()


def current_user() -> str:
    return str(st.session_state.get("username", "VISPL"))


@st.cache_data(ttl=60, show_spinner=False)
def load_people(category: str):
    response = (
        supabase.table("dropdown_master")
        .select("option_value")
        .eq("category", category)
        .eq("is_active", True)
        .order("option_value")
        .execute()
    )
    return [
        str(row.get("option_value", "")).strip()
        for row in (response.data or [])
        if str(row.get("option_value", "")).strip()
    ]


@st.cache_data(ttl=60, show_spinner=False)
def load_expense_categories():
    response = (
        supabase.table("bank_expense_categories")
        .select("category_name")
        .eq("is_active", True)
        .order("category_name")
        .execute()
    )
    return [
        str(row.get("category_name", "")).strip()
        for row in (response.data or [])
        if str(row.get("category_name", "")).strip()
    ]


def add_expense_category(category_name: str):
    clean_name = " ".join(str(category_name or "").strip().split())
    if not clean_name:
        raise ValueError("Expense category name required")
    existing = (
        supabase.table("bank_expense_categories")
        .select("id,is_active")
        .ilike("category_name", clean_name)
        .limit(1)
        .execute()
    )
    if existing.data:
        if not existing.data[0].get("is_active", True):
            supabase.table("bank_expense_categories").update(
                {"is_active": True}
            ).eq("id", existing.data[0]["id"]).execute()
            load_expense_categories.clear()
            return
        raise ValueError("यह expense category पहले से मौजूद है")
    supabase.table("bank_expense_categories").insert(
        {
            "category_name": clean_name,
            "created_by": current_user(),
            "is_active": True,
        }
    ).execute()
    load_expense_categories.clear()


def assignment_options():
    teams = load_people("Team Name")
    vendors = load_people("Vendor Name")
    return ["— Select Team/Vendor —", "Suspense", "Other Expense"] + [f"Team — {name}" for name in teams] + [f"Vendor — {name}" for name in vendors]


def parse_assignment(value: str):
    if value == "Suspense":
        return "Suspense", "Suspense"
    if value == "Other Expense":
        return "Other Expense", "Other Expense"
    if " — " not in value:
        raise ValueError("Please select a Team or Vendor")
    mode, name = value.split(" — ", 1)
    if mode not in {"Team", "Vendor"} or not name.strip():
        raise ValueError("Invalid Team/Vendor selection")
    return mode, name.strip()


def clean_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def to_text(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_amount(value):
    if pd.isna(value) or str(value).strip() == "":
        return None
    cleaned = re.sub(r"[^0-9.\-]", "", str(value).replace(",", ""))
    if cleaned in {"", "-", ".", "-."}:
        return None
    try:
        amount = round(float(cleaned), 2)
        return amount if amount > 0 else None
    except ValueError:
        return None


def to_date(value):
    if pd.isna(value) or str(value).strip() == "":
        return None
    parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def infer_pay_type(narration: str) -> str:
    text = narration.upper()
    if "PHONEPE" in text or "PHONE PE" in text or "PAYMENT FROM PHONE" in text:
        return "Phone pe"
    if "UPI" in text or "GPAY" in text or "GOOGLE PAY" in text:
        return "GPay"
    return "NEFT"


def workbook_sheet_names(file_bytes: bytes):
    return pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names


def parse_statement(file_bytes: bytes, sheet_name: str, account_label: str, file_name: str):
    raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=sheet_name,
        header=None,
        dtype=object,
    )

    required = {
        "date": "date",
        "narration": "narration",
        "chqrefno": "reference",
        "withdrawalamt": "withdrawal",
    }
    header_row = None
    column_map = {}

    for row_index in range(len(raw)):
        found = {}
        for column_index, value in enumerate(raw.iloc[row_index].tolist()):
            normalized = clean_header(value)
            if normalized in required:
                found[required[normalized]] = column_index
        if all(name in found for name in required.values()):
            header_row = row_index
            column_map = found
            break

    if header_row is None:
        raise ValueError(
            "Date, Narration, Chq./Ref.No. और Withdrawal Amt. वाली header row नहीं मिली।"
        )

    account = ACCOUNTS[account_label]
    records = []
    for row_index in range(header_row + 1, len(raw)):
        row = raw.iloc[row_index]
        txn_date = to_date(row.iloc[column_map["date"]])
        amount = to_amount(row.iloc[column_map["withdrawal"]])
        narration = to_text(row.iloc[column_map["narration"]])
        reference = to_text(row.iloc[column_map["reference"]])

        if not txn_date or amount is None or not narration:
            continue

        fingerprint_source = "|".join(
            [
                account["key"],
                txn_date.isoformat(),
                narration.strip().upper(),
                reference.strip().upper(),
                f"{amount:.2f}",
            ]
        )
        transaction_hash = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()

        records.append(
            {
                "workspace": ALLOWED_WORKSPACE,
                "account_key": account["key"],
                "pay_from": account["pay_from"],
                "transaction_date": txn_date.isoformat(),
                "narration": narration,
                "reference_no": reference,
                "withdrawal_amount": amount,
                "pay_type": infer_pay_type(narration),
                "transaction_hash": transaction_hash,
                "source_file_name": file_name,
                "source_sheet_name": sheet_name,
                "status": "Pending",
                "imported_by": current_user(),
            }
        )

    if not records:
        raise ValueError("Selected sheet में कोई Withdrawal transaction नहीं मिला।")
    return records


def parse_material_statement(file_bytes: bytes, sheet_name: str, vendor_name: str, file_name: str):
    raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=sheet_name,
        header=None,
        dtype=object,
    )

    header_aliases = {
        "invoicedate": "date",
        "date": "date",
        "billdate": "date",
        "invoiceno": "invoice_no",
        "invoicenumber": "invoice_no",
        "billno": "invoice_no",
        "billnumber": "invoice_no",
        "referenceno": "invoice_no",
        "invoiceamount": "amount",
        "billamount": "amount",
        "amount": "amount",
    }
    header_row = None
    column_map = {}
    for row_index in range(len(raw)):
        found = {}
        for column_index, value in enumerate(raw.iloc[row_index].tolist()):
            mapped_name = header_aliases.get(clean_header(value))
            if mapped_name and mapped_name not in found:
                found[mapped_name] = column_index
        if all(name in found for name in ("date", "invoice_no", "amount")):
            header_row = row_index
            column_map = found
            break

    if header_row is None:
        raise ValueError("Invoice Date, Invoice No. और Invoice Amount वाली header row नहीं मिली।")

    vendor = MATERIAL_VENDORS[vendor_name]
    records = []
    seen_hashes = set()
    for row_index in range(header_row + 1, len(raw)):
        row = raw.iloc[row_index]
        invoice_date = to_date(row.iloc[column_map["date"]])
        invoice_no = to_text(row.iloc[column_map["invoice_no"]])
        invoice_amount = to_amount(row.iloc[column_map["amount"]])
        if not invoice_date or not invoice_no or invoice_amount is None:
            continue

        fingerprint_source = "|".join(
            [
                vendor["key"],
                invoice_date.isoformat(),
                invoice_no.strip().upper(),
                f"{invoice_amount:.2f}",
            ]
        )
        transaction_hash = hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest()
        if transaction_hash in seen_hashes:
            continue
        seen_hashes.add(transaction_hash)
        material_remark = (
            f"Material Payment - Ref No. {invoice_no.strip()} "
            f"Dt. {invoice_date.strftime('%d-%b-%Y')}"
        )

        records.append(
            {
                "workspace": ALLOWED_WORKSPACE,
                "account_key": vendor["key"],
                "pay_from": vendor["pay_from"],
                "transaction_date": invoice_date.isoformat(),
                "narration": material_remark,
                "reference_no": invoice_no,
                "withdrawal_amount": invoice_amount,
                "pay_type": "Material",
                "payment_remark": material_remark,
                "transaction_hash": transaction_hash,
                "source_file_name": file_name,
                "source_sheet_name": sheet_name,
                "status": "Pending",
                "imported_by": current_user(),
            }
        )

    if not records:
        raise ValueError("Selected sheet में कोई valid material invoice नहीं मिला।")
    return records


def build_manual_material_record(vendor_name: str, invoice_date, invoice_no, invoice_amount, allow_duplicate=False):
    vendor = MATERIAL_VENDORS[vendor_name]
    clean_invoice_no = to_text(invoice_no).strip()
    clean_amount = to_amount(invoice_amount)
    if not clean_invoice_no:
        raise ValueError("Invoice Number डालें।")
    if clean_amount is None or clean_amount <= 0:
        raise ValueError("Invoice Amount 0 से ज्यादा डालें।")

    parsed_date = to_date(invoice_date)
    if not parsed_date:
        raise ValueError("Valid Invoice Date डालें।")

    fingerprint_source = "|".join(
        [
            vendor["key"],
            parsed_date.isoformat(),
            clean_invoice_no.upper(),
            f"{clean_amount:.2f}",
        ]
    )
    if allow_duplicate:
        fingerprint_source += f"|MANUAL-PROCEED|{datetime.datetime.now().isoformat()}"

    material_remark = (
        f"Material Payment - Ref No. {clean_invoice_no} "
        f"Dt. {parsed_date.strftime('%d-%b-%Y')}"
    )

    return {
        "workspace": ALLOWED_WORKSPACE,
        "account_key": vendor["key"],
        "pay_from": vendor["pay_from"],
        "transaction_date": parsed_date.isoformat(),
        "narration": material_remark,
        "reference_no": clean_invoice_no,
        "withdrawal_amount": clean_amount,
        "pay_type": "Material",
        "payment_remark": material_remark,
        "transaction_hash": hashlib.sha256(fingerprint_source.encode("utf-8")).hexdigest(),
        "source_file_name": "Manual Entry",
        "source_sheet_name": "Manual",
        "status": "Pending",
        "imported_by": current_user(),
    }


def find_manual_material_duplicate(record: dict):
    response = (
        supabase.table("bank_transactions")
        .select("id,transaction_date,reference_no,withdrawal_amount,status,pay_from")
        .eq("workspace", ALLOWED_WORKSPACE)
        .eq("account_key", record["account_key"])
        .eq("transaction_date", record["transaction_date"])
        .eq("reference_no", record["reference_no"])
        .eq("withdrawal_amount", record["withdrawal_amount"])
        .order("id", desc=True)
        .execute()
    )
    return response.data or []


def fetch_transactions(account_key: str, status: str):
    query = (
        supabase.table("bank_transactions")
        .select("*")
        .eq("workspace", ALLOWED_WORKSPACE)
        .eq("account_key", account_key)
        .eq("status", status)
        .order("transaction_date", desc=True)
        .order("id", desc=True)
    )
    return query.execute().data or []


def refresh_material_payment_details(transaction_ids):
    clean_ids = [int(value) for value in transaction_ids]
    if not clean_ids:
        return
    rows = (
        supabase.table("bank_transactions")
        .select("id,pay_type,transaction_date,reference_no")
        .in_("id", clean_ids)
        .execute()
    ).data or []
    for row in rows:
        if str(row.get("pay_type", "")).strip().lower() != "material":
            continue
        invoice_date = to_date(row.get("transaction_date"))
        invoice_no = to_text(row.get("reference_no")).strip()
        if not invoice_date or not invoice_no:
            continue
        material_remark = (
            f"Material Payment - Ref No. {invoice_no} "
            f"Dt. {invoice_date.strftime('%d-%b-%Y')}"
        )
        (
            supabase.table("bank_transactions")
            .update({
                "narration": material_remark,
                "payment_remark": material_remark,
            })
            .eq("id", int(row["id"]))
            .execute()
        )


def approve_transaction(transaction_id: int, assignment: str):
    mode, pay_to = parse_assignment(assignment)
    refresh_material_payment_details([transaction_id])
    return supabase.rpc(
        "approve_bank_transaction",
        {
            "p_transaction_id": int(transaction_id),
            "p_mode": mode,
            "p_pay_to": pay_to,
            "p_approved_by": current_user(),
        },
    ).execute()


def move_to_suspense(transaction_id: int):
    return supabase.rpc(
        "mark_bank_transaction_suspense",
        {
            "p_transaction_id": int(transaction_id),
            "p_updated_by": current_user(),
        },
    ).execute()


def mark_other_expense(transaction_id: int, expense_category: str):
    return supabase.rpc(
        "mark_bank_transaction_other_expense",
        {
            "p_transaction_id": int(transaction_id),
            "p_expense_category": expense_category,
            "p_updated_by": current_user(),
        },
    ).execute()


def restore_to_pending(transaction_id: int):
    return supabase.rpc(
        "restore_bank_transaction_pending",
        {
            "p_transaction_id": int(transaction_id),
            "p_updated_by": current_user(),
        },
    ).execute()


def bulk_move_transactions(transaction_ids, destination: str):
    if destination == "Suspense":
        destination_type = "Suspense"
        expense_category = None
    elif destination.startswith("Other Expense — "):
        destination_type = "Other Expense"
        expense_category = destination.split(" — ", 1)[1].strip()
    else:
        raise ValueError("Valid destination select करें")

    return supabase.rpc(
        "bulk_move_bank_transactions",
        {
            "p_transaction_ids": [int(value) for value in transaction_ids],
            "p_destination": destination_type,
            "p_expense_category": expense_category,
            "p_updated_by": current_user(),
        },
    ).execute()


def bulk_approve_transactions(transaction_ids, assignment: str, duplicate_action="proceed"):
    mode, pay_to = parse_assignment(assignment)
    if mode not in {"Team", "Vendor"}:
        raise ValueError("Team या Vendor select करें")
    refresh_material_payment_details(transaction_ids)
    return supabase.rpc(
        "bulk_approve_bank_transactions",
        {
            "p_transaction_ids": [int(value) for value in transaction_ids],
            "p_mode": mode,
            "p_pay_to": pay_to,
            "p_duplicate_action": duplicate_action,
            "p_approved_by": current_user(),
        },
    ).execute()


def find_possible_duplicates(row: dict, assignment: str):
    mode, pay_to = parse_assignment(assignment)
    if mode == "Suspense":
        return []
    response = (
        supabase.table("billing_payments")
        .select("id,date,amount,mode,pay_to,pay_from,pay_type,remark,workspace")
        .eq("workspace", ALLOWED_WORKSPACE)
        .eq("date", str(row.get("transaction_date")))
        .eq("amount", row.get("withdrawal_amount"))
        .eq("mode", mode)
        .eq("pay_to", pay_to)
        .order("id", desc=True)
        .execute()
    )
    return response.data or []


def link_as_duplicate(transaction_id: int, existing_payment_id: int, assignment: str):
    mode, pay_to = parse_assignment(assignment)
    return supabase.rpc(
        "link_bank_transaction_duplicate",
        {
            "p_transaction_id": int(transaction_id),
            "p_existing_payment_id": int(existing_payment_id),
            "p_mode": mode,
            "p_pay_to": pay_to,
            "p_approved_by": current_user(),
        },
    ).execute()


def format_date(value) -> str:
    parsed = pd.to_datetime(value, errors="coerce")
    return parsed.strftime("%d-%b-%Y") if not pd.isna(parsed) else ""


def format_amount(value) -> str:
    try:
        formatted = f"{float(value):,.2f}".rstrip("0").rstrip(".")
        return f"₹{formatted}"
    except (TypeError, ValueError):
        return "₹0"


def filter_transaction_records(records, search_text):
    terms = [term.lower() for term in str(search_text).split() if term.strip()]
    if not terms:
        return records
    search_fields = (
        "transaction_date", "narration", "reference_no", "withdrawal_amount",
        "status", "assignment_mode", "pay_to", "expense_category", "pay_from", "pay_type",
    )
    filtered = []
    for row in records:
        searchable_text = " ".join(str(row.get(field, "")) for field in search_fields).lower()
        if all(term in searchable_text for term in terms):
            filtered.append(row)
    return filtered


def transaction_excel(records, sheet_name):
    export_rows = []
    for row in records:
        parsed_date = pd.to_datetime(row.get("transaction_date"), errors="coerce")
        export_rows.append(
            {
                "Date": None if pd.isna(parsed_date) else parsed_date.date(),
                "Narration": str(row.get("narration", "")),
                "Chq./Ref.No.": str(row.get("reference_no", "")),
                "Withdrawal Amount": float(row.get("withdrawal_amount") or 0),
                "Status": str(row.get("status", "")),
                "Type": str(row.get("assignment_mode", "")),
                "Team/Vendor": str(row.get("pay_to", "")),
                "Expense Category": str(row.get("expense_category", "")),
                "Payment From": str(row.get("pay_from", "")),
                "Payment Type": str(row.get("pay_type", "")),
            }
        )

    export_df = pd.DataFrame(export_rows)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        worksheet = writer.sheets[sheet_name[:31]]
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="4F46E5")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, width in enumerate([15, 70, 25, 20, 18, 16, 28, 28, 25, 18], start=1):
            worksheet.column_dimensions[worksheet.cell(1, column_index).column_letter].width = width
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, 1).number_format = "DD-MMM-YYYY"
            worksheet.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.cell(row_index, 4).number_format = "#,##0.##"
    return output.getvalue()


@st.cache_data(show_spinner=False)
def material_upload_template_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Material Upload"
    worksheet.sheet_view.showGridLines = False

    worksheet.merge_cells("A1:C1")
    worksheet["A1"] = "Material Invoice Upload"
    worksheet["A1"].fill = PatternFill("solid", fgColor="4F46E5")
    worksheet["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet.merge_cells("A2:C2")
    worksheet["A2"] = "Enter one invoice per row. Do not change the three column names."
    worksheet["A2"].fill = PatternFill("solid", fgColor="EEF2FF")
    worksheet["A2"].font = Font(name="Arial", size=10, italic=True, color="3730A3")
    worksheet["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["Invoice Date", "Invoice No.", "Invoice Amount"]
    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=4, column=column_number, value=header)
        cell.fill = PatternFill("solid", fgColor="312E81")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 25
    worksheet.column_dimensions["C"].width = 20
    worksheet.freeze_panes = "A5"
    for row_number in range(5, 105):
        worksheet.cell(row_number, 1).number_format = "DD-MMM-YYYY"
        worksheet.cell(row_number, 2).number_format = "@"
        worksheet.cell(row_number, 3).number_format = "#,##0.00"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_table_toolbar(records, account_key, view_name):
    safe_view = re.sub(r"[^a-z0-9]+", "_", view_name.lower()).strip("_")
    safe_account = re.sub(r"[^a-z0-9]+", "_", account_key.lower()).strip("_")
    search_column, download_column = st.columns([4.5, 1.5])
    with search_column:
        search_text = st.text_input(
            f"Search {view_name}",
            placeholder="🔍 Search date, narration, reference, amount, team/vendor...",
            key=f"table_search_{safe_view}_{safe_account}",
            label_visibility="collapsed",
        )
    filtered_records = filter_transaction_records(records, search_text)
    with download_column:
        st.download_button(
            "📥 Download Excel",
            data=transaction_excel(filtered_records, view_name),
            file_name=f"{safe_account}_{safe_view}_{datetime.date.today():%Y%m%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"table_download_{safe_view}_{safe_account}",
            use_container_width=True,
        )
    st.caption(f"Showing {len(filtered_records)} of {len(records)} transactions")
    return filtered_records


def render_statement_preview(preview_df: pd.DataFrame):
    date_column = "Invoice Date" if "Invoice Date" in preview_df.columns else "Date"
    reference_column = "Invoice No." if "Invoice No." in preview_df.columns else "Chq./Ref.No."
    amount_column = "Invoice Amount" if "Invoice Amount" in preview_df.columns else "Withdrawal Amt."
    rows_html = []
    for _, row in preview_df.iterrows():
        rows_html.append(
            "<tr>"
            f"<td class='p-date'>{html.escape(str(row[date_column]))}</td>"
            f"<td class='p-narration'>{html.escape(str(row['Narration']))}</td>"
            f"<td class='p-ref'>{html.escape(str(row[reference_column]))}</td>"
            f"<td class='p-amount'>{html.escape(str(row[amount_column]))}</td>"
            "</tr>"
        )

    table_html = (
        "<div class='preview-wrap'><table class='preview-table'>"
        "<colgroup>"
        "<col style='width:125px'>"
        "<col>"
        "<col style='width:205px'>"
        "<col style='width:125px'>"
        "</colgroup>"
        "<thead><tr>"
        f"<th class='p-date'>{date_column}</th>"
        "<th class='p-narration'>Narration</th>"
        f"<th class='p-ref'>{reference_column}</th>"
        f"<th class='p-amount'>{amount_column}</th>"
        "</tr></thead>"
        f"<tbody>{''.join(rows_html)}</tbody>"
        "</table></div>"
    )
    st.markdown(table_html, unsafe_allow_html=True)


def render_bulk_search_and_move(records, account_key):
    with st.container(key="bulk_move_panel", border=True):
        st.markdown(
            "<div class='bulk-panel-title'>🔎 Bulk Transaction Move</div>",
            unsafe_allow_html=True,
        )
        search_column, amount_column = st.columns([3.2, 1.2])
        with search_column:
            search_text = st.text_input(
                "Search Transaction Type",
                placeholder="Search...  UPI-LITE / RENT / GST / Team name",
                key=f"bulk_search_{account_key}",
            ).strip()
        with amount_column:
            amount_less_than = st.number_input(
                "Amount Less Than",
                min_value=0,
                value=0,
                step=100,
                key=f"bulk_amount_less_than_{account_key}",
                help="Example: 1000 डालने पर केवल ₹1,000 से कम entries दिखेंगी।",
            )

        if not search_text and amount_less_than <= 0:
            return

        search_lower = search_text.lower()
        matches = [
            row
            for row in records
            if (
                not search_lower
                or search_lower in str(row.get("narration", "")).lower()
                or search_lower in str(row.get("reference_no", "")).lower()
            )
            and (
                amount_less_than <= 0
                or float(row.get("withdrawal_amount") or 0) < amount_less_than
            )
        ]

        if not matches:
            st.warning("इस search की कोई Pending entry नहीं मिली।")
            return

        total_amount = sum(float(row.get("withdrawal_amount") or 0) for row in matches)
        st.success(f"{len(matches)} matching entries मिलीं | Total {format_amount(total_amount)}")

        select_all = st.checkbox(
            "Select All",
            key=(
                f"bulk_select_all_{account_key}_"
                f"{hashlib.sha1(f'{search_lower}|{amount_less_than}'.encode()).hexdigest()[:10]}"
            ),
        )

        selection_df = pd.DataFrame(
            [
                {
                    "Select": select_all,
                    "ID": int(row["id"]),
                    "Date": format_date(row.get("transaction_date")),
                    "Narration": str(row.get("narration", "")),
                    "Chq./Ref.No.": str(row.get("reference_no", "")),
                    "Amount": format_amount(row.get("withdrawal_amount")),
                }
                for row in matches
            ]
        )

        editor_key = hashlib.sha1(
            f"{account_key}|{search_lower}|{amount_less_than}|{select_all}".encode("utf-8")
        ).hexdigest()[:12]
        edited_df = st.data_editor(
            selection_df,
            key=f"bulk_editor_{editor_key}",
            use_container_width=True,
            hide_index=True,
            disabled=["ID", "Date", "Narration", "Chq./Ref.No.", "Amount"],
            column_config={
                "Select": st.column_config.CheckboxColumn("Select", width="small"),
                "ID": None,
                "Date": st.column_config.TextColumn(width="small"),
                "Narration": st.column_config.TextColumn(width="large"),
                "Chq./Ref.No.": st.column_config.TextColumn(width="medium"),
                "Amount": st.column_config.TextColumn(width="small"),
            },
            height=min(520, 42 + (len(selection_df) * 36)),
        )

        selected_ids = edited_df.loc[edited_df["Select"] == True, "ID"].astype(int).tolist()
        expense_categories = load_expense_categories()
        destination_options = (
            ["— Select Destination —", "Suspense"]
            + [f"Other Expense — {category}" for category in expense_categories]
            + [f"Team — {name}" for name in load_people("Team Name")]
            + [f"Vendor — {name}" for name in load_people("Vendor Name")]
        )

        destination_column, action_column = st.columns([3, 1.4])
        with destination_column:
            destination = st.selectbox(
                "Move To",
                options=destination_options,
                key=f"bulk_destination_{account_key}",
            )
        with action_column:
            st.markdown("<div style='height:29px'></div>", unsafe_allow_html=True)
            move_clicked = st.button(
                f"Move Selected ({len(selected_ids)})",
                key=f"bulk_move_{account_key}",
                type="primary",
                use_container_width=True,
            )

        if move_clicked:
            try:
                if not selected_ids:
                    raise ValueError("कम से कम एक entry select करें")
                if destination == "— Select Destination —":
                    raise ValueError("Move To destination select करें")
                if destination.startswith(("Team — ", "Vendor — ")):
                    selected_id_set = set(selected_ids)
                    selected_rows = [
                        row for row in matches if int(row["id"]) in selected_id_set
                    ]
                    duplicate_rows = []
                    for selected_row in selected_rows:
                        existing_matches = find_possible_duplicates(selected_row, destination)
                        if existing_matches:
                            duplicate_rows.append(
                                {
                                    "transaction": selected_row,
                                    "existing": existing_matches,
                                }
                            )

                    if duplicate_rows:
                        st.session_state["banking_bulk_duplicate_review"] = {
                            "transaction_ids": selected_ids,
                            "assignment": destination,
                            "duplicate_rows": duplicate_rows,
                        }
                        bulk_duplicate_payment_dialog()
                        return

                    bulk_approve_transactions(selected_ids, destination, "proceed")
                    st.success(f"{len(selected_ids)} payments approve हो गईं।")
                else:
                    bulk_move_transactions(selected_ids, destination)
                    st.success(f"{len(selected_ids)} selected entries successfully move हो गईं।")
                st.rerun()
            except Exception as exc:
                st.error(f"Bulk move failed: {exc}")


def render_header(show_assignment=True):
    widths = [0.9, 5.1, 1.5, 1.1, 2.3, 1.0] if show_assignment else [0.9, 5.4, 1.5, 1.1, 2.3]
    labels = ["Date", "Narration", "Chq./Ref.No.", "Withdrawal", "Team/Vendor", "Action"] if show_assignment else ["Date", "Narration", "Chq./Ref.No.", "Withdrawal", "Booked To"]
    columns = st.columns(widths)
    for column, label in zip(columns, labels):
        column.markdown(f"<div class='table-head'>{label}</div>", unsafe_allow_html=True)


def get_paginated_records(records, state_key):
    total_pages = max(1, math.ceil(len(records) / PAGE_SIZE))
    if state_key not in st.session_state:
        st.session_state[state_key] = 1
    st.session_state[state_key] = max(
        1,
        min(int(st.session_state[state_key]), total_pages),
    )
    current_page = int(st.session_state[state_key])
    start = (current_page - 1) * PAGE_SIZE
    return records[start : start + PAGE_SIZE], current_page, total_pages


def render_pagination_buttons(state_key, current_page, total_pages):
    st.markdown("<br>", unsafe_allow_html=True)
    previous_column, count_column, next_column = st.columns([1.5, 2, 1.5])

    with previous_column:
        if st.button(
            "⬅️ Previous Page",
            key=f"previous_{state_key}",
            use_container_width=True,
            disabled=current_page <= 1,
        ):
            st.session_state[state_key] = current_page - 1
            st.rerun()

    with count_column:
        st.markdown(
            f"<div style='text-align:center;font-size:1.1rem;font-weight:800;"
            f"color:#334155;padding:13px 8px;'>Page {current_page} of {total_pages}</div>",
            unsafe_allow_html=True,
        )

    with next_column:
        if st.button(
            "Next Page ➡️",
            key=f"next_{state_key}",
            use_container_width=True,
            disabled=current_page >= total_pages,
        ):
            st.session_state[state_key] = current_page + 1
            st.rerun()


def render_pending(records, account_key, view_status="Pending"):
    if not records:
        empty_message = "Suspense में कोई transaction नहीं है।" if view_status == "Suspense" else "सभी imported transactions process हो चुके हैं।"
        st.success(empty_message)
        return

    records = render_table_toolbar(records, account_key, view_status)
    if not records:
        st.warning("Search से कोई matching transaction नहीं मिला।")
        return

    options = assignment_options()
    if not options:
        st.error("Active Team/Vendor master खाली है।")
        return

    page_state_key = f"pending_page_{view_status}_{account_key}"
    visible, current_page, page_count = get_paginated_records(records, page_state_key)

    render_header(show_assignment=True)

    for row in visible:
        row_id = int(row["id"])
        cols = st.columns([0.9, 5.1, 1.5, 1.1, 2.3, 1.0])
        cols[0].markdown(f"<div class='txn-row'><b>{format_date(row.get('transaction_date'))}</b></div>", unsafe_allow_html=True)
        safe_narration = html.escape(str(row.get("narration", "")))
        safe_reference = html.escape(str(row.get("reference_no", "")))
        cols[1].markdown(f"<div class='txn-row narration'>{safe_narration}</div>", unsafe_allow_html=True)
        cols[2].markdown(f"<div class='txn-row narration'>{safe_reference}</div>", unsafe_allow_html=True)
        cols[3].markdown(f"<div class='txn-row amount'>{format_amount(row.get('withdrawal_amount'))}</div>", unsafe_allow_html=True)
        assignment = cols[4].selectbox(
            "Team/Vendor",
            options=options,
            key=f"assignment_{view_status}_{account_key}_{row_id}",
            label_visibility="collapsed",
        )
        expense_category = None
        if assignment == "Other Expense":
            expense_categories = load_expense_categories()
            if expense_categories:
                expense_category = cols[4].selectbox(
                    "Expense Category",
                    options=expense_categories,
                    key=f"expense_category_{view_status}_{account_key}_{row_id}",
                    label_visibility="collapsed",
                )
            else:
                cols[4].error("पहले Expense Category Master में category add करें।")
        if cols[5].button("Approve", key=f"approve_{view_status}_{account_key}_{row_id}", type="primary", use_container_width=True):
            try:
                mode, _ = parse_assignment(assignment)
                if mode == "Suspense":
                    move_to_suspense(row_id)
                    st.success("Transaction Suspense में रख दिया गया।")
                    st.rerun()

                if mode == "Other Expense":
                    if not expense_category:
                        raise ValueError("Expense category select करें")
                    mark_other_expense(row_id, expense_category)
                    st.success("Transaction Other Expenses में save हुआ। Team/Vendor payment नहीं बनी।")
                    st.rerun()

                matches = find_possible_duplicates(row, assignment)
                if matches:
                    st.session_state["banking_duplicate_review"] = {
                        "transaction": row,
                        "assignment": assignment,
                        "matches": matches,
                    }
                    duplicate_payment_dialog()
                    return

                approve_transaction(row_id, assignment)
                st.success("Payment approved and booked successfully.")
                st.rerun()
            except Exception as exc:
                st.error(f"Approval failed: {exc}")

    render_pagination_buttons(page_state_key, current_page, page_count)


def render_approved(records, account_key):
    if not records:
        st.info("अभी कोई approved transaction नहीं है।")
        return

    records = render_table_toolbar(records, account_key, "Approved Payments")
    if not records:
        st.warning("Search से कोई matching transaction नहीं मिला।")
        return

    page_state_key = f"approved_page_{account_key}"
    visible, current_page, page_count = get_paginated_records(records, page_state_key)

    render_header(show_assignment=False)
    for row in visible:
        cols = st.columns([0.9, 5.4, 1.5, 1.1, 2.3])
        cols[0].markdown(f"<div class='txn-row'><b>{format_date(row.get('transaction_date'))}</b></div>", unsafe_allow_html=True)
        safe_narration = html.escape(str(row.get("narration", "")))
        safe_reference = html.escape(str(row.get("reference_no", "")))
        cols[1].markdown(f"<div class='txn-row narration'>{safe_narration}</div>", unsafe_allow_html=True)
        cols[2].markdown(f"<div class='txn-row narration'>{safe_reference}</div>", unsafe_allow_html=True)
        cols[3].markdown(f"<div class='txn-row amount'>{format_amount(row.get('withdrawal_amount'))}</div>", unsafe_allow_html=True)
        booked_to = html.escape(f"{row.get('assignment_mode', '')}: {row.get('pay_to', '')}")
        cols[4].markdown(f"<div class='txn-row approved'>{booked_to}</div>", unsafe_allow_html=True)

    render_pagination_buttons(page_state_key, current_page, page_count)


def render_other_expenses(records, account_key):
    if not records:
        st.info("अभी कोई Other Expense transaction नहीं है।")
        return

    records = render_table_toolbar(records, account_key, "Other Expenses")
    if not records:
        st.warning("Search से कोई matching transaction नहीं मिला।")
        return

    page_state_key = f"other_expense_page_{account_key}"
    visible, current_page, page_count = get_paginated_records(records, page_state_key)
    columns = st.columns([0.9, 5.0, 1.5, 1.1, 2.0, 1.2])
    for column, label in zip(
        columns,
        ["Date", "Narration", "Chq./Ref.No.", "Withdrawal", "Expense Category", "Action"],
    ):
        column.markdown(f"<div class='table-head'>{label}</div>", unsafe_allow_html=True)

    for row in visible:
        row_id = int(row["id"])
        cols = st.columns([0.9, 5.0, 1.5, 1.1, 2.0, 1.2])
        safe_narration = html.escape(str(row.get("narration", "")))
        safe_reference = html.escape(str(row.get("reference_no", "")))
        safe_category = html.escape(str(row.get("expense_category", "")))
        cols[0].markdown(f"<div class='txn-row'><b>{format_date(row.get('transaction_date'))}</b></div>", unsafe_allow_html=True)
        cols[1].markdown(f"<div class='txn-row narration'>{safe_narration}</div>", unsafe_allow_html=True)
        cols[2].markdown(f"<div class='txn-row narration'>{safe_reference}</div>", unsafe_allow_html=True)
        cols[3].markdown(f"<div class='txn-row amount'>{format_amount(row.get('withdrawal_amount'))}</div>", unsafe_allow_html=True)
        cols[4].markdown(f"<div class='txn-row approved'>{safe_category}</div>", unsafe_allow_html=True)
        if cols[5].button("Restore", key=f"restore_expense_{account_key}_{row_id}", use_container_width=True):
            try:
                restore_to_pending(row_id)
                st.success("Transaction वापस Assign & Approve में आ गया।")
                st.rerun()
            except Exception as exc:
                st.error(f"Restore failed: {exc}")

    render_pagination_buttons(page_state_key, current_page, page_count)


@st.dialog("Possible Duplicate Payment", width="large")
def duplicate_payment_dialog():
    review = st.session_state.get("banking_duplicate_review")
    if not review:
        return

    transaction = review["transaction"]
    assignment = review["assignment"]
    matches = review["matches"]

    st.warning("Same Date + Same Amount + Same Team/Vendor की payment पहले से मौजूद है।")
    st.markdown(
        f"**Current bank transaction:** {format_date(transaction.get('transaction_date'))} | "
        f"{format_amount(transaction.get('withdrawal_amount'))} | {html.escape(assignment)}"
    )

    match_df = pd.DataFrame(matches)
    show_columns = ["id", "date", "amount", "mode", "pay_to", "pay_from", "pay_type", "remark"]
    match_df = match_df[[column for column in show_columns if column in match_df.columns]].copy()
    if "date" in match_df.columns:
        match_df["date"] = match_df["date"].apply(format_date)
    if "amount" in match_df.columns:
        match_df["amount"] = match_df["amount"].apply(format_amount)
    st.dataframe(match_df, use_container_width=True, hide_index=True)

    payment_ids = [int(item["id"]) for item in matches]
    selected_payment_id = st.selectbox(
        "Existing Payment",
        options=payment_ids,
        format_func=lambda payment_id: next(
            (
                f"ID {payment_id} | {format_date(item.get('date'))} | "
                f"{format_amount(item.get('amount'))} | {item.get('mode')}: {item.get('pay_to')}"
                for item in matches
                if int(item["id"]) == payment_id
            ),
            str(payment_id),
        ),
    )

    left, right = st.columns(2)
    if left.button("Duplicate Entry", type="secondary", use_container_width=True):
        try:
            link_as_duplicate(
                int(transaction["id"]),
                int(selected_payment_id),
                assignment,
            )
            st.session_state.pop("banking_duplicate_review", None)
            st.success("Existing payment से link किया। नई payment entry नहीं बनी।")
            st.rerun()
        except Exception as exc:
            st.error(f"Duplicate link failed: {exc}")

    if right.button("Proceed", type="primary", use_container_width=True):
        try:
            approve_transaction(int(transaction["id"]), assignment)
            st.session_state.pop("banking_duplicate_review", None)
            st.success("नई payment entry save हो गई।")
            st.rerun()
        except Exception as exc:
            st.error(f"Payment approval failed: {exc}")


@st.dialog("Bulk Possible Duplicate Payments", width="large")
def bulk_duplicate_payment_dialog():
    review = st.session_state.get("banking_bulk_duplicate_review")
    if not review:
        return

    assignment = review["assignment"]
    duplicate_rows = review["duplicate_rows"]
    st.warning(
        f"{len(duplicate_rows)} selected transaction(s) में Same Date + Amount + "
        f"{html.escape(assignment)} payment पहले से मौजूद है।"
    )

    display_rows = []
    for item in duplicate_rows:
        transaction = item["transaction"]
        for existing in item["existing"]:
            display_rows.append(
                {
                    "Bank Date": format_date(transaction.get("transaction_date")),
                    "Bank Amount": format_amount(transaction.get("withdrawal_amount")),
                    "Selected For": assignment,
                    "Existing Payment ID": existing.get("id"),
                    "Existing Date": format_date(existing.get("date")),
                    "Existing Amount": format_amount(existing.get("amount")),
                    "Existing Name": f"{existing.get('mode')}: {existing.get('pay_to')}",
                }
            )
    st.dataframe(pd.DataFrame(display_rows), use_container_width=True, hide_index=True)

    left, right = st.columns(2)
    if left.button("Duplicate Entry", key="bulk_duplicate_link", use_container_width=True):
        try:
            bulk_approve_transactions(
                review["transaction_ids"],
                assignment,
                "link",
            )
            st.session_state.pop("banking_bulk_duplicate_review", None)
            st.success("Existing matching payments link हुईं; duplicate payments नहीं बनीं।")
            st.rerun()
        except Exception as exc:
            st.error(f"Bulk duplicate link failed: {exc}")

    if right.button("Proceed", key="bulk_duplicate_proceed", type="primary", use_container_width=True):
        try:
            bulk_approve_transactions(
                review["transaction_ids"],
                assignment,
                "proceed",
            )
            st.session_state.pop("banking_bulk_duplicate_review", None)
            st.success("सभी selected transactions की नई payments save हो गईं।")
            st.rerun()
        except Exception as exc:
            st.error(f"Bulk approval failed: {exc}")


@st.dialog("➕ Manual Material Invoice", width="large")
def manual_material_invoice_dialog(vendor_name: str):
    vendor = MATERIAL_VENDORS[vendor_name]
    st.markdown(f"### {html.escape(vendor['pay_from'])}")

    review_key = f"banking_manual_material_review_{vendor['key']}"
    review = st.session_state.get(review_key)

    if review:
        record = review["record"]
        duplicates = review["duplicates"]
        st.warning("यह invoice पहले से मौजूद हो सकता है। नीचे existing entry check करें।")
        duplicate_df = pd.DataFrame(
            [
                {
                    "Vendor": row.get("pay_from", vendor["pay_from"]),
                    "Invoice Date": format_date(row.get("transaction_date")),
                    "Invoice No.": row.get("reference_no", ""),
                    "Invoice Amount": format_amount(row.get("withdrawal_amount")),
                    "Status": row.get("status", ""),
                }
                for row in duplicates
            ]
        )
        st.dataframe(duplicate_df, use_container_width=True, hide_index=True)
        left, right = st.columns(2)
        if left.button("Duplicate Entry — Do Not Add", use_container_width=True):
            st.session_state.pop(review_key, None)
            st.rerun()
        if right.button("Proceed Anyway", type="primary", use_container_width=True):
            try:
                duplicate_record = build_manual_material_record(
                    vendor_name,
                    record["transaction_date"],
                    record["reference_no"],
                    record["withdrawal_amount"],
                    allow_duplicate=True,
                )
                supabase.table("bank_transactions").insert(duplicate_record).execute()
                st.session_state.pop(review_key, None)
                st.success("Manual invoice Pending में add हो गया।")
                st.rerun()
            except Exception as exc:
                st.error(f"Manual invoice save नहीं हुआ: {exc}")
        return

    with st.form(f"manual_material_form_{vendor['key']}", clear_on_submit=False):
        date_column, invoice_column, amount_column = st.columns([1.1, 1.5, 1.1])
        with date_column:
            invoice_date = st.date_input("Invoice Date", value=datetime.date.today())
        with invoice_column:
            invoice_no = st.text_input("Invoice Number", placeholder="Example: 1254/26-27")
        with amount_column:
            invoice_amount = st.number_input("Invoice Amount", min_value=0.0, step=1.0, format="%.2f")
        save_clicked = st.form_submit_button(
            "Save Manual Invoice",
            type="primary",
            use_container_width=True,
        )

    if save_clicked:
        try:
            record = build_manual_material_record(
                vendor_name,
                invoice_date,
                invoice_no,
                invoice_amount,
            )
            duplicates = find_manual_material_duplicate(record)
            if duplicates:
                st.session_state[review_key] = {
                    "record": record,
                    "duplicates": duplicates,
                }
                st.rerun(scope="fragment")
            else:
                supabase.table("bank_transactions").insert(record).execute()
                st.success("Manual invoice Pending में add हो गया।")
                st.rerun()
        except Exception as exc:
            st.error(f"Manual invoice save नहीं हुआ: {exc}")


if "banking_active_account" not in st.session_state:
    st.session_state.banking_active_account = "HDFC VISPL"
if "banking_active_material_vendor" not in st.session_state:
    st.session_state.banking_active_material_vendor = next(iter(MATERIAL_VENDORS))
if "banking_active_view" not in st.session_state:
    st.session_state.banking_active_view = "Assign & Approve"

with st.container(key="banking_account_nav"):
    account_nav_columns = st.columns(len(MAIN_ACCOUNT_TABS))
    for nav_column, account_name in zip(account_nav_columns, MAIN_ACCOUNT_TABS):
        with nav_column:
            account_active = st.session_state.banking_active_account == account_name
            if st.button(
                account_name,
                key=f"bank_account_nav_{clean_header(account_name)}",
                type="primary" if account_active else "secondary",
                use_container_width=True,
            ):
                st.session_state.banking_active_account = account_name
                st.rerun()

is_material = st.session_state.banking_active_account == "Material"
if is_material:
    with st.container(key="banking_material_vendor_nav"):
        vendor_columns = st.columns(len(MATERIAL_VENDORS))
        for vendor_column, vendor_name in zip(vendor_columns, MATERIAL_VENDORS.keys()):
            with vendor_column:
                vendor_active = st.session_state.banking_active_material_vendor == vendor_name
                if st.button(
                    vendor_name,
                    key=f"material_vendor_nav_{MATERIAL_VENDORS[vendor_name]['key']}",
                    type="primary" if vendor_active else "secondary",
                    use_container_width=True,
                ):
                    st.session_state.banking_active_material_vendor = vendor_name
                    st.rerun()
    active_account_label = st.session_state.banking_active_material_vendor
    active_account_data = MATERIAL_VENDORS[active_account_label]
    active_account_heading = f"Material — {html.escape(active_account_label)}"
else:
    active_account_label = st.session_state.banking_active_account
    active_account_data = ACCOUNTS[active_account_label]
    active_account_heading = html.escape(active_account_label)

st.markdown(
    f"""
    <div class="bank-title">
        <h1>🏦 {active_account_heading}</h1>
        <p>{'Material invoice upload, allocation and payment approval' if is_material else 'Statement upload, Team/Vendor allocation and payment approval'}</p>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.expander("Expense Category Master — नया खर्च जोड़ें"):
    st.caption("यहाँ जो category add करेंगे, वही Other Expense dropdown में दिखेगी।")
    with st.form("add_expense_category_form", clear_on_submit=True):
        new_expense_category = st.text_input(
            "New Expense Category",
            placeholder="Example: Office Rent, EMI, GST Paid",
        )
        add_category_clicked = st.form_submit_button(
            "Add Category",
            type="primary",
        )
    if add_category_clicked:
        try:
            add_expense_category(new_expense_category)
            st.success("Expense category add हो गई।")
            st.rerun()
        except Exception as exc:
            st.error(str(exc))

    try:
        active_expenses = load_expense_categories()
        if active_expenses:
            st.write("Active Categories: " + ", ".join(active_expenses))
    except Exception as exc:
        st.warning(f"Expense categories load नहीं हुईं: {exc}")

for account_label, account in [(active_account_label, active_account_data)]:
    with st.container():
        st.subheader(account_label)
        workflow_views = [
            ("Assign & Approve", "✅ Assign & Approve"),
            ("Approved Payments", "💳 Approved Payments"),
            ("Suspense", "⏳ Suspense"),
            ("Other Expenses", "🧾 Other Expenses"),
        ]
        with st.container(key="banking_view_nav"):
            view_columns = st.columns(len(workflow_views))
            for view_column, (view_key, view_label) in zip(view_columns, workflow_views):
                with view_column:
                    view_active = st.session_state.banking_active_view == view_key
                    if st.button(
                        view_label,
                        key=f"bank_view_nav_{view_key}",
                        type="primary" if view_active else "secondary",
                        use_container_width=True,
                    ):
                        st.session_state.banking_active_view = view_key
                        st.rerun()

        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown("#### Step 1: Material Invoice Upload" if is_material else "#### Step 1: Statement Upload")
        if is_material:
            manual_column, template_column, blank_column = st.columns([1.5, 1.9, 3.6])
            with manual_column:
                if st.button(
                    "➕ Manual Add Entry",
                    key=f"manual_material_entry_{account['key']}",
                    type="primary",
                    use_container_width=True,
                ):
                    manual_material_invoice_dialog(account_label)
            with template_column:
                st.download_button(
                    "⬇️ Download Material Excel Format",
                    data=material_upload_template_excel(),
                    file_name="Material_Statement_Upload_Template.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"download_material_template_{account['key']}",
                    use_container_width=True,
                )
        uploaded = st.file_uploader(
            "Upload material statement (.xls or .xlsx)" if is_material else "Upload bank statement (.xls or .xlsx)",
            type=["xls", "xlsx"],
            key=f"upload_{account['key']}",
        )

        if uploaded is not None:
            try:
                file_bytes = uploaded.getvalue()
                sheet_names = workbook_sheet_names(file_bytes)
                selected_sheet = st.selectbox(
                    "Excel Sheet",
                    options=sheet_names,
                    key=f"sheet_{account['key']}",
                )
                if is_material:
                    preview_records = parse_material_statement(
                        file_bytes,
                        selected_sheet,
                        account_label,
                        uploaded.name,
                    )
                    st.info(f"Preview: {len(preview_records)} material invoices मिले। Existing duplicate invoices import नहीं होंगे।")
                else:
                    preview_records = parse_statement(
                        file_bytes,
                        selected_sheet,
                        account_label,
                        uploaded.name,
                    )
                    st.info(f"Preview: {len(preview_records)} withdrawal transactions मिले। Deposit और Closing Balance शामिल नहीं हैं।")

                preview_df = pd.DataFrame(preview_records)[
                    ["transaction_date", "narration", "reference_no", "withdrawal_amount"]
                ].sort_values("transaction_date", ascending=False)
                preview_df.columns = [
                    "Invoice Date" if is_material else "Date",
                    "Narration",
                    "Invoice No." if is_material else "Chq./Ref.No.",
                    "Invoice Amount" if is_material else "Withdrawal Amt.",
                ]
                date_column = "Invoice Date" if is_material else "Date"
                amount_column = "Invoice Amount" if is_material else "Withdrawal Amt."
                preview_df[date_column] = pd.to_datetime(preview_df[date_column]).dt.strftime("%d-%b-%Y")
                preview_df[amount_column] = preview_df[amount_column].apply(format_amount)
                render_statement_preview(preview_df)

                if st.button("Save Transactions & Continue", key=f"import_{account['key']}", type="primary"):
                    before = sum(
                        len(fetch_transactions(account["key"], status))
                        for status in ("Pending", "Approved", "Suspense", "Other Expense")
                    )
                    supabase.table("bank_transactions").upsert(
                        preview_records,
                        on_conflict="workspace,account_key,transaction_hash",
                        ignore_duplicates=True,
                    ).execute()
                    after = sum(
                        len(fetch_transactions(account["key"], status))
                        for status in ("Pending", "Approved", "Suspense", "Other Expense")
                    )
                    st.success(f"Import completed. New entries: {max(0, after - before)} | Duplicate skipped: {max(0, len(preview_records) - max(0, after - before))}")
                    st.rerun()
            except Exception as exc:
                st.error(f"Statement save नहीं हुआ: {exc}")

        if st.session_state.banking_active_view == "Assign & Approve":
            try:
                pending_records = fetch_transactions(account["key"], "Pending")
                render_bulk_search_and_move(pending_records, account["key"])
                render_pending(pending_records, account["key"], "Pending")
            except Exception as exc:
                st.warning(f"Supabase connect नहीं हुआ, इसलिए Team/Vendor dropdown अभी नहीं दिख सकता: {exc}")
        elif st.session_state.banking_active_view == "Approved Payments":
            try:
                render_approved(fetch_transactions(account["key"], "Approved"), account["key"])
            except Exception as exc:
                st.warning(f"Approved payments load नहीं हुए: {exc}")
        elif st.session_state.banking_active_view == "Suspense":
            try:
                render_pending(fetch_transactions(account["key"], "Suspense"), account["key"], "Suspense")
            except Exception as exc:
                st.warning(f"Suspense transactions load नहीं हुए: {exc}")
        elif st.session_state.banking_active_view == "Other Expenses":
            try:
                render_other_expenses(
                    fetch_transactions(account["key"], "Other Expense"),
                    account["key"],
                )
            except Exception as exc:
                st.warning(f"Other Expenses load नहीं हुए: {exc}")
